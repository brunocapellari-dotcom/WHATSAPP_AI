"""Report export regressions: bytes, totals and untrusted spreadsheet input."""
from io import BytesIO
from datetime import datetime
from decimal import Decimal

from openpyxl import load_workbook

from app.freight_reporting import FreightObservation, FreightReportEngine, MT_TIMEZONE, simplify
from app.report_exports import (
    full_xlsx,
    report_pdf,
    safe_spreadsheet_text,
    simplified_xlsx,
)
def dt(day: int, hour: int, minute: int = 0) -> datetime:
    return datetime(2026, 10, day, hour, minute, tzinfo=MT_TIMEZONE)


def offer(source: str, carrier: str = "A", loading_point: str = "Fazenda") -> FreightObservation:
    return FreightObservation(source_observation_id=source, published_at=dt(15, 7),
        carrier=carrier, origin_city="Sorriso", origin_uf="MT",
        destination_city="Rondonópolis", destination_uf="MT", product="SOJA",
        loading_point=loading_point, rate=Decimal("180"), rate_unit="BRL/t")


def test_formula_injection_is_neutralized() -> None:
    for prefix in ("=", "+", "-", "@", "\t"):
        assert safe_spreadsheet_text(prefix + "SUM(1,1)").startswith("'")
    assert safe_spreadsheet_text("Safe") == "Safe"


def test_simplified_export_has_two_main_columns_and_totals() -> None:
    engine = FreightReportEngine()
    full = engine.process(dt(15, 8, 30), [
        offer("export-a", carrier="A", loading_point="=HYPERLINK(1)"),
        offer("export-b", carrier="B", loading_point="=HYPERLINK(1)"),
    ])
    simple = simplify(full)
    wb = load_workbook(BytesIO(simplified_xlsx(simple)), read_only=True)
    ws = wb.active
    assert [ws["A1"].value, ws["B1"].value] == ["Origem / Destino", "Frete"]
    assert ws["A2"].data_type == "s"
    assert any(row[0] == "Total de lotes abertos" and row[1] == 1
               for row in ws.values)


def test_full_export_and_pdf_have_valid_file_signatures() -> None:
    engine = FreightReportEngine()
    full = engine.process(dt(15, 8, 30), [offer("export-1")])
    assert full_xlsx(full).startswith(b"PK")
    assert simplified_xlsx(simplify(full)).startswith(b"PK")
    assert report_pdf(full).startswith(b"%PDF")
    assert report_pdf(simplify(full)).startswith(b"%PDF")
